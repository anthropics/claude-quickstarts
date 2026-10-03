"""Offline tests for the notebook's document normalization functions."""

import ast
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches

NOTEBOOK = (
    Path(__file__).resolve().parents[1] / "distill_documents_into_knowledge_wiki.ipynb"
)


def load_normalizers():
    notebook = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    names = {"normalize_docx", "normalize_any"}
    namespace = {"Path": Path}
    for cell in notebook["cells"]:
        if cell["cell_type"] != "code":
            continue
        source = "".join(cell["source"])
        if not any(f"def {name}(" in source for name in names):
            continue
        definitions = [
            node
            for node in ast.parse(source).body
            if isinstance(node, ast.FunctionDef) and node.name in names
        ]
        module = ast.Module(body=definitions, type_ignores=[])
        # Load only these definitions, without executing notebook setup or API cells.
        exec(compile(module, str(NOTEBOOK), "exec"), namespace)  # noqa: S102
    return {name: namespace[name] for name in names}


class NormalizationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.normalizers = load_normalizers()

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.directory = Path(self.temp_dir.name)

    def test_docx_keeps_tables_between_their_sections(self):
        path = self.directory / "sections.docx"
        document = Document()
        document.add_heading("Quarter 1", level=1)
        first = document.add_table(rows=2, cols=2)
        first.cell(0, 0).text = "Metric"
        first.cell(0, 1).text = "Value"
        first.cell(1, 0).text = "Q1 revenue"
        first.cell(1, 1).text = "100"
        document.add_heading("Quarter 2", level=1)
        second = document.add_table(rows=2, cols=2)
        second.cell(0, 0).text = "Metric"
        second.cell(0, 1).text = "Value"
        second.cell(1, 0).text = "Q2 revenue"
        second.cell(1, 1).text = "200"
        document.add_paragraph("End of report")
        document.save(path)

        self.assertEqual(
            self.normalizers["normalize_docx"](path),
            "# Quarter 1\n\n| Metric | Value |\n\n|---|---|\n\n"
            "| Q1 revenue | 100 |\n\n# Quarter 2\n\n| Metric | Value |\n\n"
            "|---|---|\n\n| Q2 revenue | 200 |\n\nEnd of report",
        )

    def test_docx_keeps_existing_terminal_table_format(self):
        path = self.directory / "terminal-table.docx"
        document = Document()
        document.add_heading("Quarter 1", level=1)
        document.add_paragraph("Summary")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Revenue"
        table.cell(1, 1).text = "100"
        document.save(path)

        self.assertEqual(
            self.normalizers["normalize_docx"](path),
            "# Quarter 1\n\nSummary\n\n| Metric | Value |\n\n|---|---|\n\n"
            "| Revenue | 100 |",
        )

    def test_pptx_includes_native_table_contents(self):
        path = self.directory / "table.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5), Inches(5), Inches(0.5)
        )
        textbox.text = "Quarterly results"
        table = slide.shapes.add_table(
            2, 2, Inches(0.5), Inches(1.2), Inches(5), Inches(1)
        ).table
        table.cell(0, 0).text = "Metric"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Revenue"
        table.cell(1, 1).text = "12345"
        presentation.save(path)

        self.assertEqual(
            self.normalizers["normalize_any"](path),
            "## Slide 1\nQuarterly results\n| Metric | Value |\n|---|---|\n"
            "| Revenue | 12345 |",
        )

    def test_pptx_keeps_existing_text_only_format(self):
        path = self.directory / "text-only.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5), Inches(5), Inches(0.5)
        )
        textbox.text = "Existing text slide"
        presentation.save(path)

        self.assertEqual(
            self.normalizers["normalize_any"](path),
            "## Slide 1\nExisting text slide",
        )


if __name__ == "__main__":
    unittest.main()
